#!/usr/bin/env python3
"""Build the Syracuse Budget Explorer data file from the Auditor workbook.

The adopted PDF supplies the formal FY27 totals and amendments. The workbook
supplies the reusable account-level history. This script joins both sources.
"""

from __future__ import annotations

import argparse
import ast
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from validate_budget import validate


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT.parent / "tmp" / "2026-city-budget-workbook.xlsx"
DEFAULT_OUTPUT = ROOT / "data" / "budget.json"

PDF_URL = (
    "https://www.syr.gov/files/sharedassets/public/v/3/departments/budget/"
    "documents/budget-documents/08-12-fy27-final-adopted-budget.pdf"
)
BUDGET_PAGE_URL = "https://www.syr.gov/Departments/Budget/City-Budget"
AUDITOR_PAGE_URL = "https://www.syr.gov/Departments/Auditor-Office"
WORKBOOK_URL = (
    "https://www.syr.gov/files/content/public/v/94/departments/"
    "auditor-office/2026-city-budget-workbook.xlsx"
)
OPEN_DATA_URL = "https://data.syr.gov/"

YEAR_KEYS = (
    "fy21Actual",
    "fy22Actual",
    "fy23Actual",
    "fy24Actual",
    "fy25Actual",
    "fy26Budget",
    "fy26Estimate",
    "fy27Proposed",
)

FUND_CONFIG = {
    "General Fund": {
        "id": "general-fund",
        "name": "General Fund",
        "kind": "operating",
        "formal": {"fy26Budget": 333036573, "fy27Adopted": 350640243},
    },
    "Water Fund": {
        "id": "water-fund",
        "name": "Water",
        "kind": "enterprise",
        "formal": {"fy26Budget": 31031846, "fy27Adopted": 33175883},
    },
    "Sewer Fund": {
        "id": "sewer-fund",
        "name": "Sewer",
        "kind": "enterprise",
        "formal": {"fy26Budget": 7483800, "fy27Adopted": 8966902},
    },
    "Sidewalk Fund": {
        "id": "sidewalk-fund",
        "name": "Municipal Sidewalk",
        "kind": "enterprise",
        "formal": {"fy26Budget": 2285237, "fy27Adopted": 2719688},
    },
    "Crouse Marshall Special": {
        "id": "crouse-marshall",
        "name": "Crouse-Marshall assessment",
        "kind": "special",
        "formal": {"fy26Budget": 176567, "fy27Adopted": 181863},
    },
    "Downtown Special": {
        "id": "downtown-assessment",
        "name": "Downtown assessment",
        "kind": "special",
        "formal": {"fy26Budget": 1370770, "fy27Adopted": 1411892},
    },
}

AMENDMENTS = [
    {
        "id": "temp-aim-aid",
        "section": "revenue",
        "label": "Temporary AIM state aid",
        "detail": "Account 435209",
        "amount": 20000000,
        "code": "435209",
    },
    {
        "id": "fund-balance",
        "section": "revenue",
        "label": "Unreserved, undesignated fund balance",
        "detail": "The adopted plan leaves a residual $879 instead of using the proposed draw.",
        "amount": -23881500,
        "code": "FUND_BALANCE",
        "adopted_value": -879,
    },
    {
        "id": "social-security",
        "section": "expense",
        "label": "Social Security",
        "detail": "Special Objects of Expense, account 590301",
        "amount": -500000,
        "code": "590301",
        "workbook_codes": ["590300"],
    },
    {
        "id": "medical-insurance",
        "section": "expense",
        "label": "Medical Insurance",
        "detail": "Special Objects of Expense, account 590601",
        "amount": -900000,
        "code": "590601",
        "workbook_codes": ["590600"],
    },
    {
        "id": "retirement-system",
        "section": "expense",
        "label": "Employee Retirement System",
        "detail": "Special Objects of Expense, account 590101",
        "amount": -900000,
        "code": "590101",
        "workbook_codes": ["590100"],
    },
    {
        "id": "financial-management-system",
        "section": "expense",
        "label": "Financial Management System",
        "detail": "Special Objects of Expense, account 599891",
        "amount": -51500,
        "code": "599891",
    },
    {
        "id": "allowance-negotiations",
        "section": "expense",
        "label": "Allowance for negotiations",
        "detail": "Special Objects of Expense, account 590051",
        "amount": -1000000,
        "code": "590051",
        "workbook_codes": ["590050"],
    },
    {
        "id": "cash-capital",
        "section": "expense",
        "label": "Cash capital appropriations",
        "detail": "Expense line Cash Capital",
        "amount": -2120000,
        "code": "CASH_CAPITAL",
        "name_terms": ["cash capital appropriations"],
    },
    {
        "id": "onondaga-historical-association",
        "section": "expense",
        "label": "Onondaga Historical Association",
        "detail": "Special Objects of Expense, account 594500",
        "amount": 25000,
        "code": "594500",
    },
    {
        "id": "public-events",
        "section": "expense",
        "label": "Public Events",
        "detail": "Special Objects of Expense, account 595500",
        "amount": 50000,
        "code": "595500",
    },
    {
        "id": "home-headquarters",
        "section": "expense",
        "label": "Home HeadQuarters",
        "detail": "Special Objects of Expense, account 595946",
        "amount": 750000,
        "code": "595946",
    },
    {
        "id": "one-time-expenditures",
        "section": "expense",
        "label": "One Time Expenditures",
        "detail": "Special Objects of Expense, account 593000",
        "amount": 100000,
        "code": "593000",
    },
    {
        "id": "parks-recreation-services",
        "section": "expense",
        "label": "Parks Recreation professional services",
        "detail": "Division of Recreation, account 541500",
        "amount": 175000,
        "code": "541500",
        "workbook_codes": ["541500"],
        "department_terms": ["parks"],
        "division_terms": ["recreation"],
        "name_terms": ["professional services"],
    },
    {
        "id": "parks-wages",
        "section": "expense",
        "label": "Parks Grounds Maintenance wages",
        "detail": "Division of Parks/Grounds Maintenance, account 510200",
        "amount": 50000,
        "code": "510200",
        "department_terms": ["parks"],
        "division_terms": ["grounds maintenance"],
        "name_terms": ["wages"],
    },
    {
        "id": "parks-facility-repair",
        "section": "expense",
        "label": "Parks Grounds Maintenance facility repair",
        "detail": "Division of Parks/Grounds Maintenance, account 540511",
        "amount": 125000,
        "code": "540511",
        "department_terms": ["parks"],
        "division_terms": ["grounds maintenance"],
        "name_terms": ["facility repair"],
    },
    {
        "id": "law-salaries",
        "section": "expense",
        "label": "Department of Law salaries",
        "detail": "Department of Law, account 510100",
        "amount": 100000,
        "code": "510100",
        "department_terms": ["department of law"],
        "name_terms": ["salaries"],
    },
    {
        "id": "personnel-salaries",
        "section": "expense",
        "label": "Personnel and Labor Relations salaries",
        "detail": "Office of Personnel and Labor Relations, account 510100",
        "amount": 80000,
        "code": "510100",
        "division_terms": ["personnel"],
        "name_terms": ["salaries"],
    },
    {
        "id": "common-council-salaries",
        "section": "expense",
        "label": "Common Council salaries",
        "detail": "Common Council, account 510100",
        "amount": 100000,
        "code": "510100",
        "department_terms": ["common council"],
        "name_terms": ["salaries"],
    },
    {
        "id": "city-clerk-services",
        "section": "expense",
        "label": "City Clerk professional services",
        "detail": "City Clerk, account 541500",
        "amount": 35000,
        "code": "541500",
        "department_terms": ["city clerk"],
        "name_terms": ["professional services"],
    },
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").replace("\n", " ")
    text = text.replace("_", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def slugify(value: str) -> str:
    text = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return text or "line"


def strip_code(value: str) -> str:
    return re.sub(r"\s*-\s*\d{4,6}\s*$", "", value).strip()


def evaluate_arithmetic(value: str) -> float | None:
    expression = value.strip().replace(",", "")
    if expression.startswith("="):
        expression = expression[1:].strip()
    if not expression:
        return None
    try:
        tree = ast.parse(expression, mode="eval").body
    except SyntaxError:
        return None
    return evaluate_node(tree)


def evaluate_node(node: ast.AST) -> float | None:
    if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
        return float(node.value)
    if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
        number = evaluate_node(node.operand)
        if number is None:
            return None
        return number if isinstance(node.op, ast.UAdd) else -number
    if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub, ast.Mult, ast.Div)):
        left = evaluate_node(node.left)
        right = evaluate_node(node.right)
        if left is None or right is None:
            return None
        if isinstance(node.op, ast.Add):
            return left + right
        if isinstance(node.op, ast.Sub):
            return left - right
        if isinstance(node.op, ast.Mult):
            return left * right
        if right == 0:
            return None
        return left / right
    return None


def number(value: Any) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 2)
    text = clean(value)
    if text in {"", "N/A", "NA", "-"}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    text = text.replace("$", "").replace("%", "").replace("−", "-")
    parsed = evaluate_arithmetic(text)
    if parsed is None:
        return None
    return round(-parsed if negative else parsed, 2)


def year_key(label: str, sheet_name: str) -> str | None:
    match = re.search(r"FY(\d{2})\s+(Actual|Budget|Estimate|Proposed)", label, re.I)
    if not match:
        return None
    year, kind = match.groups()
    if sheet_name == "Debt Service" and year == "25" and kind.lower() == "budget":
        year = "26"
    return f"fy{year}{kind.title()}"


def find_year_columns(ws, sheet_name: str) -> dict[str, int]:
    columns: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 3)):
        for cell in row:
            key = year_key(clean(cell.value), sheet_name)
            if key and key not in columns:
                columns[key] = cell.column
    return columns


def account_match(value: str) -> re.Match[str] | None:
    return re.search(r"(?<!\d)(\d{6})\s+(.+)$", value)


def section_for(text: str) -> str | None:
    lower = text.lower()
    if lower == "revenues" or lower.startswith("revenues "):
        return "revenue"
    if lower == "expenses" or lower.startswith("expenses "):
        return "expense"
    return None


def is_total_label(text: str) -> bool:
    lower = text.lower()
    return (
        lower.startswith("subtotal")
        or lower.startswith("total ")
        or lower.startswith("grand total")
        or lower in {"revenues", "expenses"}
    )


def line_label(row: list[Any], label_limit: int) -> tuple[str, int]:
    choices: list[tuple[str, int]] = []
    for index in range(label_limit):
        text = clean(row[index])
        if text and text != "`":
            choices.append((text, index + 1))
    if not choices:
        return "", 0
    return choices[-1]


def context_name(contexts: dict[int, str], column: int) -> str | None:
    value = contexts.get(column)
    if not value:
        return None
    return strip_code(value)


def row_values(row: list[Any], year_columns: dict[str, int]) -> dict[str, float]:
    values: dict[str, float] = {}
    for key, column in year_columns.items():
        parsed = number(row[column - 1]) if column <= len(row) else None
        if parsed is not None:
            values[key] = parsed
    return values


def should_keep_line(sheet_name: str, label: str) -> bool:
    # Debt service, reserves and transfers often have no account code.
    # Keep numeric leaf lines; hierarchy filtering below excludes their subtotals.
    return bool(label) and not is_total_label(label)


def record_from_row(
    *,
    fund_id: str,
    fund_name: str,
    sheet_name: str,
    section: str,
    row_number: int,
    row: list[Any],
    contexts: dict[int, str],
    year_columns: dict[str, int],
    label_limit: int,
    account: re.Match[str] | None,
    account_column: int,
) -> dict[str, Any] | None:
    values = row_values(row, year_columns)
    if not values:
        return None
    label = clean(account.group(2)) if account else line_label(row, label_limit)[0]
    if not label or is_total_label(label):
        return None
    code = account.group(1) if account else None
    prefix = clean(account.group(0)[: account.group(0).find(account.group(1))]) if account else ""
    if prefix:
        label = clean(account.group(0)[len(prefix) :])
    if not account and not should_keep_line(sheet_name, label):
        return None

    if account:
        level = "account"
        if sheet_name in {"Water Fund", "Sewer Fund", "Sidewalk Fund"} and account_column <= 2:
            category = "Revenue" if section == "revenue" else "Expense"
        elif sheet_name == "General Fund" and section == "revenue":
            category = prefix or context_name(contexts, 2)
        else:
            category_column = max(2, account_column - 1)
            category = prefix or context_name(contexts, category_column)
        department = context_name(contexts, 3) if account_column >= 4 else None
        division = context_name(contexts, 4) if account_column >= 5 else None
    else:
        level = "line"
        category = context_name(contexts, 2)
        department = context_name(contexts, 3)
        division = context_name(contexts, 4)

    if sheet_name in {"Crouse Marshall Special", "Downtown Special"}:
        department = fund_name

    row_id_parts = [fund_id, section, code or label, str(row_number)]
    return {
        "id": slugify("-".join(row_id_parts)),
        "fundId": fund_id,
        "fundName": fund_name,
        "section": section,
        "level": level,
        "category": category,
        "department": department,
        "division": division,
        "code": code,
        "name": label,
        "values": values,
        "source": "2026 City Budget Workbook",
        "sourceSheet": sheet_name,
        "sourceRow": row_number,
        "adoptedMethod": (
            "workbook_detail_plus_adopted_amendments"
            if sheet_name == "General Fund"
            else "workbook_detail_formal_fund_total"
        ),
    }


def parse_sheet(ws, sheet_name: str) -> list[dict[str, Any]]:
    config = FUND_CONFIG.get(sheet_name)
    if not config:
        return []
    year_columns = find_year_columns(ws, sheet_name)
    if not year_columns:
        return []
    first_year_column = min(year_columns.values())
    label_limit = first_year_column - 1
    contexts: dict[int, str] = {}
    current_section: str | None = None
    special_assessment = False
    records: list[dict[str, Any]] = []

    for row_number in range(1, ws.max_row + 1):
        row = [ws.cell(row_number, column).value for column in range(1, ws.max_column + 1)]
        labels = [clean(value) for value in row[:label_limit]]
        row_section = next((section_for(label) for label in labels if section_for(label)), None)
        if row_section:
            current_section = row_section
            special_assessment = False

        account: re.Match[str] | None = None
        account_column = 0
        for index, label in enumerate(labels):
            match = account_match(label)
            if match:
                account = match
                account_column = index + 1
                break

        line_name, line_column = line_label(row, label_limit)
        if line_name.lower().startswith("special assessment levy"):
            special_assessment = True
        if special_assessment and not row_section:
            effective_section = "revenue"
        else:
            effective_section = current_section

        if effective_section and (account or should_keep_line(sheet_name, line_name)):
            record = record_from_row(
                fund_id=config["id"],
                fund_name=config["name"],
                sheet_name=sheet_name,
                section=effective_section,
                row_number=row_number,
                row=row,
                contexts=contexts,
                year_columns=year_columns,
                label_limit=label_limit,
                account=account,
                account_column=account_column or line_column,
            )
            # Indentation defines hierarchy in this workbook. A row followed by
            # a deeper label is a subtotal/group, even when its label contains
            # an account code (for example Contractual Expenses + Office Supplies).
            is_parent = False
            current_column = account_column or line_column
            for next_number in range(row_number + 1, ws.max_row + 1):
                next_labels = [(col, clean(ws.cell(next_number, col).value)) for col in range(1, label_limit + 1)]
                next_labels = [(col, value) for col, value in next_labels if value and value != "`"]
                if next_labels:
                    is_parent = next_labels[-1][0] > current_column
                    break
            if record and not is_parent:
                records.append(record)

        for index, label in enumerate(labels, start=1):
            match = account_match(label)
            if match:
                prefix = clean(label[: match.start()])
                if prefix:
                    contexts[index] = prefix
                for key in list(contexts):
                    if key > index and not prefix:
                        contexts.pop(key, None)
                continue
            if label and label != "`" and not is_total_label(label):
                contexts[index] = label
                for key in list(contexts):
                    if key > index:
                        contexts.pop(key, None)

    return records


def normalise(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()


def amendment_codes(amendment: dict[str, Any]) -> set[str]:
    return {amendment["code"], *amendment.get("workbook_codes", [])}


def match_score(row: dict[str, Any], amendment: dict[str, Any]) -> int:
    if row["fundId"] != "general-fund" or row["section"] != amendment["section"]:
        return -1
    terms = amendment.get("name_terms", [])
    row_name = normalise(row.get("name"))
    if terms and not all(normalise(term) in row_name for term in terms):
        return -1
    score = 0
    if row.get("code") in amendment_codes(amendment):
        score += 10
    for field, key in (("department", "department_terms"), ("division", "division_terms")):
        field_value = normalise(row.get(field))
        if amendment.get(key) and not all(normalise(term) in field_value for term in amendment[key]):
            return -1
        if amendment.get(key):
            score += 5
    if amendment.get("code") == "FUND_BALANCE" and row.get("name", "").lower().startswith("unreserved"):
        score += 20
    if amendment.get("code") == "CASH_CAPITAL" and "cash capital" in row_name:
        score += 20
    return score


def apply_amendments(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    applied: list[dict[str, Any]] = []
    for amendment in AMENDMENTS:
        candidates = [(match_score(row, amendment), row) for row in rows]
        candidates = [(score, row) for score, row in candidates if score >= 10]
        if not candidates:
            row = {
                "id": slugify(f"general-fund-{amendment['section']}-{amendment['code']}"),
                "fundId": "general-fund",
                "fundName": "General Fund",
                "section": amendment["section"],
                "level": "line",
                "category": "Adopted amendment",
                "department": None,
                "division": None,
                "code": amendment["code"],
                "name": amendment["label"],
                "values": {"fy27Proposed": 0, "fy27Adjusted": amendment["amount"]},
                "source": "FY2026–27 adopted budget PDF",
                "sourceSheet": "Subsequent Events",
                "sourceRow": None,
                "adoptedMethod": "adopted_pdf_amendment_only",
                "note": "This amendment line was not present as a separate workbook row.",
            }
            rows.append(row)
            applied.append({"id": amendment["id"], "rowId": row["id"], "amount": amendment["amount"]})
            continue
        candidates.sort(key=lambda item: item[0], reverse=True)
        row = candidates[0][1]
        proposed = row["values"].get("fy27Proposed", 0)
        adopted = amendment.get("adopted_value", proposed + amendment["amount"])
        row["values"]["fy27Adjusted"] = round(float(adopted), 2)
        row.setdefault("amendments", []).append(amendment["id"])
        applied.append(
            {
                "id": amendment["id"],
                "rowId": row["id"],
                "amount": amendment["amount"],
            }
        )
    return applied


def change_for(values: dict[str, float], base: str = "fy26Budget", target: str = "fy27Adjusted") -> dict[str, float | None]:
    old = values.get(base, 0) or 0
    new = values.get(target, values.get("fy27Proposed", 0)) or 0
    change = round(new - old, 2)
    percent = None if old == 0 else round(change / abs(old) * 100, 1)
    return {"amount": change, "pct": percent}


def add_changes(rows: list[dict[str, Any]]) -> None:
    for row in rows:
        row["change"] = change_for(row["values"])


def formal_item(name: str, old: int, new: int) -> dict[str, Any]:
    change = new - old
    return {
        "name": name,
        "fy26Budget": old,
        "fy27Adopted": new,
        "change": change,
        "pct": round(change / old * 100, 1) if old else None,
    }


def build_summary() -> dict[str, Any]:
    city_old = 372044043
    city_new = 393755721
    combined_old = 983045670
    combined_new = 1027852721
    return {
        "city": formal_item("City funds, net", city_old, city_new),
        "generalFund": formal_item("General Fund", 333036573, 350640243),
        "schoolDistrict": formal_item("Syracuse City School District", 611001627, 634097000),
        "otherCityFunds": formal_item("All other City funds", 42348220, 46456228),
        "combinedNet": formal_item("City and school district, net", combined_old, combined_new),
        "interfund": -3340750,
        "revenueSources": [
            {"label": "State aid, net of STAR", "amount": 620324666, "pct": 60.35, "scope": "combined"},
            {"label": "Real property taxes", "amount": 130359277, "pct": 12.68, "scope": "combined"},
            {"label": "Non-property taxes", "amount": 131086341, "pct": 12.75, "scope": "combined"},
            {"label": "Other revenues", "amount": 136355630, "pct": 13.27, "scope": "combined"},
            {"label": "Real property tax items", "amount": 9726807, "pct": 0.95, "scope": "combined"},
        ],
        "spendingSources": [
            {"label": "School district", "amount": 634097000, "pct": 61.7, "scope": "combined"},
            {"label": "Employee benefits", "amount": 114844549, "pct": 11.2, "scope": "combined"},
            {"label": "Police", "amount": 63746688, "pct": 6.2, "scope": "combined"},
            {"label": "Fire", "amount": 47895330, "pct": 4.7, "scope": "combined"},
            {"label": "Public Works", "amount": 41121923, "pct": 4.0, "scope": "combined"},
            {"label": "Principal and interest", "amount": 26495345, "pct": 2.6, "scope": "combined"},
            {"label": "Parks and Recreation", "amount": 12547384, "pct": 1.2, "scope": "combined"},
            {"label": "Other City operations and capital", "amount": 87104501, "pct": 8.5, "scope": "combined"},
        ],
        "narrative": [
            {
                "label": "The adopted plan adds aid",
                "body": "Temporary AIM state aid rises by $20.0M from the proposed workbook plan. The adopted PDF leaves a $879 fund-balance residual.",
                "kind": "revenue",
            },
            {
                "label": "The General Fund grows",
                "body": "The adopted General Fund is $350.6M, up $17.6M, or 5.3%, from the FY26 adopted plan.",
                "kind": "change",
            },
            {
                "label": "Final amendments matter",
                "body": "The final General Fund is $3.9M below the workbook's FY27 proposed plan after the listed amendments.",
                "kind": "amendment",
            },
        ],
    }


def build_funds() -> list[dict[str, Any]]:
    funds: list[dict[str, Any]] = []
    for config in FUND_CONFIG.values():
        formal = config["formal"]
        old = formal["fy26Budget"]
        new = formal["fy27Adopted"]
        funds.append(
            {
                "id": config["id"],
                "name": config["name"],
                "kind": config["kind"],
                "formal": formal_item(config["name"], old, new),
            }
        )
    return funds


def build_data(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    workbook_values = load_workbook(workbook_path, data_only=True, read_only=False)
    rows: list[dict[str, Any]] = []
    formula_cells = 0
    formula_errors = 0
    for sheet_name in FUND_CONFIG:
        ws = workbook[sheet_name]
        value_ws = workbook_values[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formula_cells += 1
                    if number(cell.value) is None and number(value_ws[cell.coordinate].value) is None:
                        formula_errors += 1
        rows.extend(parse_sheet(value_ws, sheet_name))

    applied = apply_amendments(rows)
    add_changes(rows)
    result = {
        "meta": {
            "title": "Syracuse Budget Explorer",
            "fiscalPeriod": "FY2026–27",
            "fiscalDates": "July 1, 2026 – June 30, 2027",
            "lastUpdated": "September 4, 2026",
            "scope": "Formal adopted City fund totals, with the school district shown separately in the combined summary.",
            "sourceNote": "The Auditor workbook supplies historical account detail. The adopted PDF supplies formal FY27 totals and final amendments.",
            "workbookCaveat": "The Auditor notes that workbook figures can differ from the formal budget because of rounding or other errors.",
        },
        "years": [
            {"key": "fy21Actual", "label": "FY21 actual"},
            {"key": "fy22Actual", "label": "FY22 actual"},
            {"key": "fy23Actual", "label": "FY23 actual"},
            {"key": "fy24Actual", "label": "FY24 actual"},
            {"key": "fy25Actual", "label": "FY25 actual"},
            {"key": "fy26Budget", "label": "FY26 adopted"},
            {"key": "fy26Estimate", "label": "FY26 estimate"},
            {"key": "fy27Proposed", "label": "FY27 proposed"},
            {"key": "fy27Adjusted", "label": "FY27 adjusted proposal"},
            {"key": "fy27Adopted", "label": "FY27 adopted (unverified)"},
        ],
        "summary": build_summary(),
        "funds": build_funds(),
        "rows": rows,
        "amendments": [
            {
                key: value
                for key, value in amendment.items()
                if key not in {"workbook_codes", "duplicate", "adopted_value", "name_terms", "department_terms", "division_terms"}
            }
            for amendment in AMENDMENTS
        ],
        "appliedAmendments": applied,
        "validation": [],
        "sources": {
            "budgetPage": {"label": "City budget page", "url": BUDGET_PAGE_URL},
            "adoptedPdf": {"label": "FY2026–27 adopted budget PDF", "url": PDF_URL},
            "auditorPage": {"label": "City Auditor budget workbook page", "url": AUDITOR_PAGE_URL},
            "workbook": {"label": "2026 City Budget Workbook XLSX", "url": WORKBOOK_URL},
            "openData": {"label": "Syracuse Open Data Portal", "url": OPEN_DATA_URL},
        },
        "openData": {
            "catalogUrl": OPEN_DATA_URL,
            "note": "Older public layers support cross-checks. FY27 adopted totals come from the formal PDF and workbook.",
            "layers": [
                {
                    "label": "Approved Budget Revenue, FY2023–24",
                    "url": "https://services6.arcgis.com/bdPqSfflsdgFRVVM/arcgis/rest/services/SYRPROD_FY24_Annual_Budget_Revenue/FeatureServer",
                },
                {
                    "label": "Proposed Budget Revenue, FY2024–25",
                    "url": "https://services6.arcgis.com/bdPqSfflsdgFRVVM/arcgis/rest/services/Proposed_Budget_Revenue_(Fiscal_Year_2024_2025)/FeatureServer",
                },
                {
                    "label": "Proposed Budget Expenditures, FY2024–25",
                    "url": "https://services6.arcgis.com/bdPqSfflsdgFRVVM/arcgis/rest/services/Proposed_Budget_Expenditures_(Fiscal_Year_2024_to_2025)/FeatureServer",
                },
            ],
        },
    }
    result["validation"] = validate(result)
    result["validation"].append({"label": "Workbook formula scan", "status": "pass" if formula_errors == 0 else "check", "detail": f"{formula_cells} formulas scanned; {formula_errors} unreadable. This checks readability, not formula correctness."})
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize Syracuse budget workbook data.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    data = build_data(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(json.dumps({"rows": len(data["rows"]), "appliedAmendments": len(data["appliedAmendments"]), "output": str(args.output)}))


if __name__ == "__main__":
    main()
