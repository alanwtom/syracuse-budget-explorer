import unittest
from openpyxl import Workbook
from ingest_budget import parse_sheet


class ParserTests(unittest.TestCase):
    def test_embedded_account_in_subtotal_is_not_a_leaf(self):
        ws = Workbook().active
        ws.append(['Budget', None, None, 'FY26 Budget', 'FY27 Proposed'])
        ws.append(['Expenses'])
        ws.append([None, 'Contractual Expenses\n540300 Office Supplies', None, 100, 600])
        ws.append([None, None, '540300 Office Supplies', 10, 20])
        ws.append([None, None, '541500 Professional Services', 90, 580])
        rows = parse_sheet(ws, 'General Fund')
        self.assertEqual([r['values']['fy27Proposed'] for r in rows], [20, 580])

    def test_unnumbered_debt_and_transfers_are_retained(self):
        ws = Workbook().active
        ws.append(['Budget', None, 'FY26 Budget', 'FY27 Proposed'])
        ws.append(['Expenses'])
        ws.append([None, 'Debt service principal', 100, 120])
        ws.append([None, 'Transfer to other funds', 20, 30])
        rows = parse_sheet(ws, 'Water Fund')
        self.assertEqual(sum(r['values']['fy27Proposed'] for r in rows), 150)

    def test_assessment_parent_is_not_counted_twice(self):
        ws = Workbook().active
        ws.append(['Budget', None, 'FY26 Budget', 'FY27 Proposed'])
        ws.append(['Special Assessment Levy', None, 100, 110])
        ws.append([None, 'Special Assessment', 100, 110])
        rows = parse_sheet(ws, 'Crouse Marshall Special')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['section'], 'revenue')


if __name__ == '__main__':
    unittest.main()
